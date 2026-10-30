import json
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from spreadsheet_exports import build_exports


def products():
    return pd.DataFrame(
        [
            {"id": 1, "nome": "Café & chá", "preco": "R$ 29,90", "ativo": True},
            {"id": 2, "nome": "Ação de cozinha", "preco": "R$ 15,00", "ativo": False},
        ]
    )


def test_exports_include_four_formats_with_correct_names_and_mime_types():
    exports = build_exports(products())

    assert set(exports) == {"CSV", "Excel", "TSV", "JSON"}
    assert exports["CSV"]["file_name"].endswith(".csv")
    assert exports["Excel"]["file_name"].endswith(".xlsx")
    assert exports["TSV"]["file_name"].endswith(".tsv")
    assert exports["JSON"]["file_name"].endswith(".json")
    assert all(isinstance(item["data"], bytes) and item["data"] for item in exports.values())


def test_csv_and_tsv_preserve_portuguese_text_and_table_columns():
    exports = build_exports(products())

    csv = pd.read_csv(BytesIO(exports["CSV"]["data"]))
    tsv = pd.read_csv(BytesIO(exports["TSV"]["data"]), sep="\t")
    for table in (csv, tsv):
        assert table.columns.tolist() == ["id", "nome", "preco", "ativo"]
        assert table.loc[0, "nome"] == "Café & chá"
        assert table.loc[1, "nome"] == "Ação de cozinha"


def test_excel_export_has_products_sheet_frozen_header_and_filter():
    exports = build_exports(products())
    workbook = load_workbook(BytesIO(exports["Excel"]["data"]), data_only=True)
    worksheet = workbook["Produtos"]

    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:D3"
    assert worksheet.cell(row=1, column=2).value == "nome"
    assert worksheet.cell(row=2, column=2).value == "Café & chá"
    assert worksheet.cell(row=3, column=4).value is False
    assert worksheet.cell(row=1, column=1).fill.fgColor.rgb.endswith("4F46E5")


def test_json_export_is_utf8_records_with_boolean_values():
    exports = build_exports(products())
    records = json.loads(exports["JSON"]["data"].decode("utf-8"))

    assert records[0] == {"id": 1, "nome": "Café & chá", "preco": "R$ 29,90", "ativo": True}
    assert records[1]["ativo"] is False
