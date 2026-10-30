"""Utilities for exporting the product table in portable spreadsheet formats."""

from __future__ import annotations

from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def build_exports(frame: pd.DataFrame) -> dict[str, dict[str, Any]]:
    """Return named, downloadable representations of a product DataFrame."""
    table = frame.copy()

    csv_data = table.to_csv(index=False).encode("utf-8-sig")
    tsv_data = table.to_csv(index=False, sep="\t").encode("utf-8-sig")
    json_data = table.to_json(orient="records", force_ascii=False, indent=2).encode("utf-8")

    xlsx_buffer = BytesIO()
    with pd.ExcelWriter(xlsx_buffer, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Produtos", index=False)
        worksheet = writer.sheets["Produtos"]
        worksheet.freeze_panes = "A2"
        if worksheet.max_column:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="4F46E5")
        for column_index, cells in enumerate(worksheet.columns, start=1):
            longest = max((len(str(cell.value)) if cell.value is not None else 0) for cell in cells)
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(longest + 2, 12), 40)
    xlsx_data = xlsx_buffer.getvalue()

    return {
        "CSV": {
            "data": csv_data,
            "file_name": "axadinhos_vippe_produtos.csv",
            "mime": "text/csv; charset=utf-8",
        },
        "Excel": {
            "data": xlsx_data,
            "file_name": "axadinhos_vippe_produtos.xlsx",
            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        },
        "TSV": {
            "data": tsv_data,
            "file_name": "axadinhos_vippe_produtos.tsv",
            "mime": "text/tab-separated-values; charset=utf-8",
        },
        "JSON": {
            "data": json_data,
            "file_name": "axadinhos_vippe_produtos.json",
            "mime": "application/json; charset=utf-8",
        },
    }
